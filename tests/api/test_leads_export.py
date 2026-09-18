"""Exporting My Leads as CSV and Excel.

Leads are seeded straight into the database: what is under test is the file,
not the search that produced the rows.
"""

import csv
import io
import uuid
from datetime import UTC, datetime
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.config import settings
from app.models.lead import Company, DecisionMaker
from app.models.user import User
from app.services.lead_export import COLUMNS
from tests.api.test_auth import signed_in_tokens

EXPORT = f"{settings.API_V1_PREFIX}/leads/export"


async def _headers(client: AsyncClient, db_session: Any) -> dict[str, str]:
    tokens = await signed_in_tokens(client, db_session)
    return {"Authorization": f"Bearer {tokens['access_token']}"}


async def _user_id(db_session: Any) -> str:
    return (await db_session.execute(select(User))).scalars().first().id


async def _lead(
    db_session: Any,
    user_id: str,
    name: str,
    *,
    status: str = "new",
    in_leads: bool = True,
    contacts: list[dict[str, Any]] | None = None,
    **fields: Any,
) -> Company:
    company = Company(
        user_id=user_id,
        company_name=name,
        website=fields.pop("website", f"{name.lower().replace(' ', '-')}.example.com"),
        status=status,
        added_to_leads_at=datetime.now(UTC) if in_leads else None,
        extra_json="{}",
        **fields,
    )
    db_session.add(company)
    await db_session.flush()

    for contact in contacts or []:
        db_session.add(DecisionMaker(company_id=company.id, extra_json="{}", **contact))

    await db_session.commit()

    return company


def _csv_rows(body: bytes) -> list[list[str]]:
    assert body.startswith(b"\xef\xbb\xbf"), "Excel needs the UTF-8 BOM"
    return list(csv.reader(io.StringIO(body.decode("utf-8-sig"))))


def _xlsx_rows(body: bytes) -> list[list[str]]:
    sheet = load_workbook(io.BytesIO(body)).active
    return [
        [("" if value is None else str(value)) for value in row]
        for row in sheet.iter_rows(values_only=True)
    ]


# --- Shape -------------------------------------------------------------------


async def test_csv_has_one_row_per_contact(client: AsyncClient, db_session) -> None:
    headers = await _headers(client, db_session)
    user_id = await _user_id(db_session)
    await _lead(
        db_session,
        user_id,
        "Acme GmbH",
        location="Berlin",
        contacts=[
            {
                "full_name": "Jane Roe",
                "job_title": "CPO",
                "verified_email": "jane@acme.example.com",
                "email_status": "valid",
                "phone_number": "+49 30 1234",
                "whatsapp_status": "Active",
            },
            {"full_name": "Tom Li", "whatsapp_status": "Not Active"},
        ],
    )

    response = await client.get(EXPORT, params={"format": "csv"}, headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert 'filename="leadpilot-leads-all-' in response.headers["content-disposition"]

    rows = _csv_rows(response.content)
    assert rows[0] == list(COLUMNS)
    assert len(rows) == 3, "a header and one row per contact"

    jane, tom = rows[1], rows[2]
    assert jane[:3] == ["Acme GmbH", "acme-gmbh.example.com", "Berlin"]
    assert jane[6] == "new"
    assert jane[9:] == [
        "Jane Roe", "CPO", "jane@acme.example.com", "valid", "", "+49 30 1234", "Active"
    ]
    # The company columns are repeated so each row stands on its own.
    assert tom[:3] == jane[:3]
    assert tom[9] == "Tom Li"
    assert tom[15] == "Not Active"


async def test_a_company_without_contacts_still_gets_a_row(
    client: AsyncClient, db_session
) -> None:
    headers = await _headers(client, db_session)
    user_id = await _user_id(db_session)
    await _lead(db_session, user_id, "Lonely Ltd")

    rows = _csv_rows((await client.get(EXPORT, headers=headers)).content)

    assert len(rows) == 2
    assert rows[1][0] == "Lonely Ltd"
    assert rows[1][9:] == [""] * 7


async def test_an_unchecked_whatsapp_number_is_a_blank_cell(
    client: AsyncClient, db_session
) -> None:
    """Null means never checked; it must not come out as 'None'."""
    headers = await _headers(client, db_session)
    user_id = await _user_id(db_session)
    await _lead(
        db_session,
        user_id,
        "Acme",
        contacts=[{"full_name": "A", "whatsapp_status": None}],
    )

    rows = _csv_rows((await client.get(EXPORT, headers=headers)).content)

    assert rows[1][15] == ""


async def test_xlsx_matches_the_csv(client: AsyncClient, db_session) -> None:
    headers = await _headers(client, db_session)
    user_id = await _user_id(db_session)
    await _lead(
        db_session,
        user_id,
        "Acme",
        contacts=[{"full_name": "Jane Roe", "verified_email": "jane@acme.example.com"}],
    )

    response = await client.get(EXPORT, params={"format": "xlsx"}, headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="leadpilot-leads-all-' in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.xlsx"')

    rows = _xlsx_rows(response.content)
    assert rows[0] == list(COLUMNS)
    assert rows[1][0] == "Acme"
    assert rows[1][9] == "Jane Roe"
    assert rows[1][11] == "jane@acme.example.com"


async def test_the_default_format_is_csv(client: AsyncClient, db_session) -> None:
    headers = await _headers(client, db_session)

    response = await client.get(EXPORT, headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")


# --- Scope -------------------------------------------------------------------


async def test_status_filter_narrows_the_file(client: AsyncClient, db_session) -> None:
    headers = await _headers(client, db_session)
    user_id = await _user_id(db_session)
    await _lead(db_session, user_id, "New Co", status="new")
    await _lead(db_session, user_id, "Warm Co", status="contacted")

    response = await client.get(
        EXPORT, params={"status": "contacted"}, headers=headers
    )

    rows = _csv_rows(response.content)
    assert [row[0] for row in rows[1:]] == ["Warm Co"]

    disposition = response.headers["content-disposition"]
    assert 'filename="leadpilot-leads-contacted-' in disposition


async def test_ids_export_just_the_selection(client: AsyncClient, db_session) -> None:
    headers = await _headers(client, db_session)
    user_id = await _user_id(db_session)
    keep = await _lead(db_session, user_id, "Keep Co")
    await _lead(db_session, user_id, "Skip Co")

    response = await client.get(EXPORT, params={"ids": [keep.id]}, headers=headers)

    rows = _csv_rows(response.content)
    assert [row[0] for row in rows[1:]] == ["Keep Co"]


async def test_an_id_that_is_not_yours_is_ignored(
    client: AsyncClient, db_session
) -> None:
    """Scoped by user: a foreign id yields nothing rather than someone's data."""
    headers = await _headers(client, db_session)
    user_id = await _user_id(db_session)
    mine = await _lead(db_session, user_id, "Mine Co")

    response = await client.get(
        EXPORT, params={"ids": [mine.id, str(uuid.uuid4())]}, headers=headers
    )

    assert [row[0] for row in _csv_rows(response.content)[1:]] == ["Mine Co"]


async def test_search_results_not_added_to_leads_are_excluded(
    client: AsyncClient, db_session
) -> None:
    headers = await _headers(client, db_session)
    user_id = await _user_id(db_session)
    await _lead(db_session, user_id, "Lead Co", in_leads=True)
    await _lead(db_session, user_id, "Result Co", in_leads=False)

    rows = _csv_rows((await client.get(EXPORT, headers=headers)).content)

    assert [row[0] for row in rows[1:]] == ["Lead Co"]


async def test_an_empty_export_is_just_the_header(
    client: AsyncClient, db_session
) -> None:
    headers = await _headers(client, db_session)

    rows = _csv_rows((await client.get(EXPORT, headers=headers)).content)

    assert rows == [list(COLUMNS)]


# --- Validation --------------------------------------------------------------


async def test_an_unknown_format_is_rejected(client: AsyncClient, db_session) -> None:
    headers = await _headers(client, db_session)

    response = await client.get(EXPORT, params={"format": "pdf"}, headers=headers)

    assert response.status_code == 422


async def test_export_requires_auth(client: AsyncClient) -> None:
    assert (await client.get(EXPORT)).status_code == 401


async def test_export_is_not_mistaken_for_a_lead_id(
    client: AsyncClient, db_session
) -> None:
    """`/leads/export` must reach the export route, not `/leads/{lead_id}`."""
    headers = await _headers(client, db_session)

    response = await client.get(EXPORT, headers=headers)

    assert response.status_code == 200, response.text

"""Exporting My Leads as a file.

One row per decision maker, with the company's columns repeated beside them, so
every contact field sits in its own column and the file drops straight into a
spreadsheet, a mail merge or a CRM import. A company with no contacts still
gets one row, so no lead goes missing from the file.

CSV is written as UTF-8 with a byte-order mark: without it Excel on Windows
reads accented names as mojibake. XLSX is built with openpyxl in memory --
exports are at most a few thousand rows, so there is nothing to stream.
"""

import csv
import io
from collections.abc import Iterable, Iterator
from datetime import datetime
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.lead import Company, DecisionMaker
from app.repositories.lead import CompanyRepository
from app.services.base import BaseService

ExportFormat = Literal["csv", "xlsx"]

CONTENT_TYPES: dict[str, str] = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

# Header labels, in column order. Company first, then the contact.
COLUMNS: tuple[str, ...] = (
    "Company",
    "Website",
    "Location",
    "Industry",
    "Company size",
    "HQ phone",
    "Status",
    "Added on",
    "Notes",
    "Contact",
    "Job title",
    "Email",
    "Email status",
    "LinkedIn",
    "Phone",
    "WhatsApp",
)

_CONTACT_COLUMNS = 7


def _text(value: object) -> str:
    return "" if value is None else str(value)


def _company_cells(company: Company) -> list[str]:
    added = company.added_to_leads_at
    return [
        company.company_name,
        _text(company.website),
        _text(company.location),
        _text(company.industry),
        _text(company.company_size),
        _text(company.hq_phone),
        company.status,
        added.strftime("%Y-%m-%d") if isinstance(added, datetime) else "",
        _text(company.notes),
    ]


def _contact_cells(person: DecisionMaker) -> list[str]:
    return [
        person.full_name,
        _text(person.job_title),
        _text(person.verified_email),
        _text(person.email_status),
        _text(person.linkedin_url),
        _text(person.phone_number),
        # Null means the number was never checked, and a blank cell says so
        # more honestly than any word would.
        _text(person.whatsapp_status),
    ]


def rows_for(companies: Iterable[Company]) -> Iterator[list[str]]:
    """The data rows, in the order the leads page lists them."""
    for company in companies:
        base = _company_cells(company)

        if not company.decision_makers:
            yield [*base, *([""] * _CONTACT_COLUMNS)]
            continue

        for person in company.decision_makers:
            yield [*base, *_contact_cells(person)]


def render_csv(companies: Iterable[Company]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    writer.writerows(rows_for(companies))

    # utf-8-sig: the BOM is what makes Excel read it as UTF-8.
    return buffer.getvalue().encode("utf-8-sig")


def render_xlsx(companies: Iterable[Company]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active

    # A fresh workbook always has one plain worksheet; the stubs only say
    # "maybe a sheet, maybe a chart sheet", so narrow it for the type checker.
    assert isinstance(sheet, Worksheet)
    sheet.title = "Leads"

    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    row_count = 0
    for row in rows_for(companies):
        sheet.append(row)
        row_count += 1

    # Header stays put while scrolling, and can be filtered.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{row_count + 1}"

    # Readable widths without measuring every cell: enough for a URL or an
    # email, narrower for short codes.
    for index, label in enumerate(COLUMNS, start=1):
        wide = label in {"Company", "Website", "Email", "LinkedIn", "Notes", "Contact"}
        sheet.column_dimensions[get_column_letter(index)].width = 32 if wide else 16

    buffer = io.BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()


def filename_for(fmt: ExportFormat, status: str | None, at: datetime) -> str:
    """leadpilot-leads-contacted-2026-09-17.xlsx, say."""
    return f"leadpilot-leads-{status or 'all'}-{at:%Y-%m-%d}.{fmt}"


class LeadExportService(BaseService):
    def __init__(self, db: AsyncSession) -> None:
        super().__init__(db)
        self.companies = CompanyRepository(db)

    async def export(
        self,
        user_id: str,
        *,
        fmt: ExportFormat,
        status: str | None = None,
        ids: list[str] | None = None,
        search: str | None = None,
    ) -> bytes:
        """The file body for this user's leads, narrowed by status, ids and/or
        the My Leads search text."""
        search = " ".join((search or "").split()) or None
        companies = await self.companies.list_leads_for_export(
            user_id, status=status, ids=ids, search=search
        )

        return render_xlsx(companies) if fmt == "xlsx" else render_csv(companies)


__all__ = [
    "COLUMNS",
    "CONTENT_TYPES",
    "ExportFormat",
    "LeadExportService",
    "filename_for",
    "render_csv",
    "render_xlsx",
    "rows_for",
]
